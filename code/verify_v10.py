"""verify_v10.py — Round-10 targeted checks after reviewer-flagged text fixes.

Confirms:
(1) 'an no self-reported diagnosis' grammar error is gone
(2) 'severe-or-proliferative' terminology consistent everywhere (no bare 'proliferative severity' /
    'proliferative disease')
(3) Table 5 note no longer claims 'conditional rate within the NHIS diabetic subsample'
(4) Table 5 note now correctly describes VIM_DREV as universal to all NHIS 40+ Sample Adults
(5) The NHIS DR denominator trace artifact exists in results/ and its 8.18% conditional value
    is documented (so a future author can re-derive if needed)
(6) Abstract first sentence softened (no 'has not been characterised')
(7) Intro 'Both are largely asymptomatic until late' phrase split for DR vs AMD
(8) 'unreported diagnosis' replaced with 'no self-reported diagnosis' / equivalent
(9) Cross-reference identifiers consistent between main docx and Supplementary docx
"""
from pathlib import Path
import re, sys
import docx
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
MD   = ROOT / "manuscript" / "题目3手稿最新版本.docx"
SUPP = ROOT / "manuscript" / "Supplementary_Materials_v7.docx"
NHIS_TRACE = ROOT / "results" / "sr_nhis_dr_denom_trace_v10.csv"

def read_text(path):
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for tbl in d.tables:
        for row in tbl.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)

oks, fails = [], []
def check(name, cond, detail=""):
    (oks if cond else fails).append((name, bool(cond), detail))

main_txt = read_text(MD)
supp_txt = read_text(SUPP) if SUPP.exists() else ""
both_txt = main_txt + "\n" + supp_txt

# (1) grammar error 'an no'
n = len(re.findall(r'an\s+no\s+self-reported', both_txt))
check("(1) no 'an no self-reported' grammar", n == 0, "%d occurrences" % n)

# (2) terminology consistency: no bare 'proliferative severity' or 'proliferative disease'
bad_prolif = re.findall(r'(?<!severe-or-)\bproliferative\s+(severity|disease)\b', both_txt, re.I)
check("(2) no bare 'proliferative severity/disease'", len(bad_prolif) == 0, "found %s" % bad_prolif[:3])
check("(2b) 'severe-or-proliferative retinopathy' present",
      'severe-or-proliferative retinopathy' in main_txt)

# (3) wrong conditional claim removed
check("(3) 'conditional rate within the NHIS diabetic subsample' removed",
      'conditional rate within the NHIS diabetic subsample' not in main_txt)

# (4) new correct claim present (VIM_DREV asked of ALL Sample Adults)
check("(4) Table 5 note now describes VIM_DREV as universal",
      ('VIM_DREV' in main_txt and 'without a diabetes branch screener' in main_txt) or
      ('asked without a diabetes branch screener' in main_txt))

# (5) NHIS trace artifact exists with 8.18% conditional
if NHIS_TRACE.exists():
    tr = pd.read_csv(NHIS_TRACE)
    c_row = tr[tr['scenario'].str.contains('C:')]
    if len(c_row):
        pct = float(c_row.iloc[0]['pct'])
        check("(5) NHIS trace file present with conditional ~8.18%",
              7.8 <= pct <= 8.6, "diabetics-only weighted %%: %.2f" % pct)
    else:
        check("(5) NHIS trace file has scenario C", False)
else:
    check("(5) NHIS trace artifact exists", False, str(NHIS_TRACE))

# (6) Abstract softened
check("(6) abstract no longer says 'has not been characterised in nationally'",
      'has not been characterised in nationally' not in main_txt)
check("(6b) abstract uses 'remains incompletely characterised'",
      'remains incompletely characterised' in main_txt)

# (7) Intro DR/AMD split
check("(7) intro differentiates DR vs AMD onset",
      'Diabetic retinopathy is usually asymptomatic' in main_txt)

# (8) 'unreported diagnosis' fully removed
n_unrep = len(re.findall(r'\bunreported diagnosis\b', both_txt))
check("(8) 'unreported diagnosis' fully removed", n_unrep == 0, "%d found" % n_unrep)

# (9) Cross-reference consistency
main_ids = set(re.findall(r'Supplementary Table (S\d+[a-z]?)', main_txt))
supp_ids = set(re.findall(r'Supplementary Table (S\d+[a-z]?)\.', supp_txt))
check("(9a) all main-text supplementary refs have supplement sections",
      main_ids <= supp_ids, "missing in supp: %s" % (main_ids - supp_ids))
check("(9b) all supplement sections cited in main text",
      supp_ids <= main_ids, "not cited: %s" % (supp_ids - main_ids))

# Summary
print("=" * 60)
print("Round-10 verify:  PASSED: %d   FAILED: %d" % (len(oks), len(fails)))
print("=" * 60)
if fails:
    print()
    print("Failures:")
    for name, cond, detail in fails:
        print("  [FAIL] %s   %s" % (name, detail))
    sys.exit(1)
else:
    print("ALL ROUND-10 CHECKS PASS")
