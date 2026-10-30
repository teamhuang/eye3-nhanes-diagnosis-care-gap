import docx, re

DOCX = 'manuscript/题目3手稿最新版本.docx'
doc = docx.Document(DOCX)

issues = []
patterns = [
    (r'an\s+no\s+', "grammar: 'an no'"),
    (r'(?<!severe-or-)\bproliferative\s+(disease|severity|category)\b', "standalone proliferative-X"),
    (r'conditional rate within.*?diabetic subsample', "wrong conditional-diabetic claim"),
]
for i, p in enumerate(doc.paragraphs):
    for pat, desc in patterns:
        for m in re.finditer(pat, p.text):
            issues.append((i, desc, p.text[max(0, m.start()-40):m.end()+40]))
    for ti, tbl in enumerate(doc.tables):
        for ri, row in enumerate(tbl.rows):
            for ci, cell in enumerate(row.cells):
                for pat, desc in patterns:
                    if re.search(pat, cell.text):
                        issues.append(('T%d[%d,%d]' % (ti, ri, ci), desc, cell.text[:80]))

if not issues:
    print('PASS: no P0 text errors remain in main docx')
else:
    for i, d, c in issues:
        print('para/table %s [%s]: %s' % (i, d, c))
