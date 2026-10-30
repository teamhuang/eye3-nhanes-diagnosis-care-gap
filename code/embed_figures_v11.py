"""Replace the 3 embedded images in the main docx with the new Nature-style PNGs,
at exact 5 cm height and their specified widths (Fig1/Fig2 = 6.5 cm, Fig3 = 5 cm)."""
import docx
from docx.shared import Cm
from docx.oxml.ns import qn
from pathlib import Path

DOCX = Path('manuscript/题目3手稿最新版本.docx')
FIGS = Path('figures')

doc = docx.Document(str(DOCX))

# Locate image paragraphs and their target file
targets = [
    (76, "Fig1_gaps.png", 6.5, 5.0),
    (78, "Fig2_severity_gradient.png", 6.5, 5.0),
    (80, "Fig3_objective_vs_selfreport.png", 5.0, 5.0),
]

for pi, fname, w_cm, h_cm in targets:
    p = doc.paragraphs[pi]
    # Remove all existing runs (which contain the old image)
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    # Add new image
    run = p.add_run()
    run.add_picture(str(FIGS/fname), width=Cm(w_cm), height=Cm(h_cm))
    print('para %d: replaced with %s (%.1f x %.1f cm)' % (pi, fname, w_cm, h_cm))

doc.save(str(DOCX))
print('Saved.')
