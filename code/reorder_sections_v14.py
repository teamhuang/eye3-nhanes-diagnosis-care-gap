"""Reorder main-manuscript sections: move the Methods block to BEFORE Results.
Target body order: Title/Abstract/Keywords -> Introduction -> Methods -> Results ->
Discussion -> Conclusions -> Declarations(Ethics..Competing) -> Figure legends -> Tables -> References.
Ethics/Code/Data/Funding/Contributions/COI stay in back matter (they are declarations, not the
scientific Methods)."""
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

ROOT = Path("C:/eye3/manuscript/eye3")
D = ROOT/"manuscript"/"题目3手稿最新版本.docx"
d = Document(str(D))
ps = d.paragraphs

def find_head(title):
    for i,p in enumerate(ps):
        if p.text.strip()==title: return i
    raise SystemExit("heading not found: "+title)

methods_h = find_head("Methods")
ethics_i  = next(i for i,p in enumerate(ps) if p.text.strip().startswith("Ethics."))
results_h = find_head("Results")

# Methods scientific block = [methods_h .. ethics_i-1]
block = [ps[i]._p for i in range(methods_h, ethics_i)]
results_el = ps[results_h]._p
print(f"methods_h={methods_h} ethics_i={ethics_i} results_h={results_h} block={len(block)} paras")

# move block (preserving order) to just before the Results heading
for el in block:
    results_el.addprevious(el)

d.save(str(D))
# re-open to report new order
d2=Document(str(D))
heads=[p.text.strip() for p in d2.paragraphs if p.style and p.style.name in ("Heading 1","Heading 2")]
print("NEW H2 ORDER:", heads)
