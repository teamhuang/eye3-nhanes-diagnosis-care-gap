#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""format_v15.py — Round-15 docx formatting.
(1) reorder: move the 'Figure legends .. Tables' display block to the very END of the main
    manuscript (after References), keeping supplementary figures/tables out of the main file.
(2) three-line (booktabs) borders on EVERY table in the main and supplementary docx:
    thick top + bottom rule, single rule under the header row, no vertical / no inner rules.
    Captions stay bold (already are); no other text is bolded.
Idempotent: re-running just re-applies borders / re-moves an already-at-end block (no-op)."""
import sys, docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def _txt(el): return ''.join(t.text or '' for t in el.iter(qn('w:t'))).strip()

def set_three_line(table):
    tblPr = table._tbl.tblPr
    tb = tblPr.find(qn('w:tblBorders'))
    if tb is None:
        tb = OxmlElement('w:tblBorders')
        # insert after tblW if present else at position 0
        tw = tblPr.find(qn('w:tblW'))
        (tw.addnext(tb) if tw is not None else tblPr.insert(0, tb))
    def edge(tag, sz, val):
        e = tb.find(qn('w:'+tag))
        if e is None:
            e = OxmlElement('w:'+tag); tb.append(e)
        e.set(qn('w:val'), val); e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), '0'); e.set(qn('w:color'), 'auto')
    edge('top',12,'single'); edge('bottom',12,'single')
    edge('left',0,'none'); edge('right',0,'none'); edge('insideH',0,'none'); edge('insideV',0,'none')
    # header row bottom rule
    for cell in table.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        for ex in tcPr.findall(qn('w:tcBorders')): tcPr.remove(ex)
        tcb = OxmlElement('w:tcBorders')
        b = OxmlElement('w:bottom'); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'6'); b.set(qn('w:space'),'0'); b.set(qn('w:color'),'auto')
        tcb.append(b)
        shd = tcPr.find(qn('w:shd'))
        (shd.addprevious(tcb) if shd is not None else tcPr.append(tcb))

def reorder_to_end(doc):
    body = doc.element.body
    els = list(body.iterchildren())
    si = next((i for i,e in enumerate(els) if e.tag==qn('w:p') and _txt(e).startswith('Figure legends')), None)
    ei = next((i for i,e in enumerate(els) if e.tag==qn('w:p') and _txt(e).startswith('Supplementary Tables')), None)
    if si is None or ei is None or si>=ei:
        print('reorder: block not found (si=%s ei=%s) — skipping'%(si,ei)); return 0
    block = els[si:ei]
    sectPr = body.find(qn('w:sectPr'))
    for e in block: sectPr.addprevious(e)
    return len(block)

def borders_all(doc):
    n=0
    for t in doc.tables:
        set_three_line(t); n+=1
    return n

if __name__=='__main__':
    main=sys.argv[1]; supp=sys.argv[2]
    d=docx.Document(main)
    moved=reorder_to_end(d)
    nb=borders_all(d)
    d.save(main)
    print('MAIN: moved %d elements to end; %d tables set to three-line'%(moved,nb))
    s=docx.Document(supp)
    nbs=borders_all(s)
    s.save(supp)
    print('SUPP: %d tables set to three-line'%nbs)
